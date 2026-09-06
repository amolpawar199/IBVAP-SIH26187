import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    # Initialize 16:9 Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Use blank slide layout
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Define Theme Colors
    BG_COLOR = RGBColor(11, 15, 25)          # #0b0f19 Dark Tech Background
    CARD_BG = RGBColor(30, 41, 59)           # #1e293b Slate Card Fill
    INNER_CARD_BG = RGBColor(15, 23, 42)    # #0f172a Sub-card Fill
    ACCENT_CYAN = RGBColor(56, 189, 248)     # #38bdf8 Cyan
    ACCENT_GREEN = RGBColor(16, 185, 129)    # #10b981 Emerald
    TEXT_WHITE = RGBColor(248, 250, 252)     # #f8fafc Text Bright
    TEXT_MUTED = RGBColor(148, 163, 184)    # #94a3b8 Text Gray

    # 1. Background Fill
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background()

    # 2. Header Title & Subtitle
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(9.5), Inches(0.9))
    tf = title_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "IBVAP – Intelligent Border Video Analytics Platform"
    p.font.name = 'Arial'
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    p2 = tf.add_paragraph()
    p2.text = "AI-Powered Perimeter Defense & Autonomous Real-Time Threat Intelligence"
    p2.font.name = 'Arial'
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_MUTED

    # Badge Component
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.3), Inches(0.45), Inches(2.4), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(14, 165, 233)
    badge.line.color.rgb = ACCENT_CYAN
    tf_b = badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "NEXT-GEN AI SECURITY"
    p_b.font.name = 'Arial'
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = RGBColor(255, 255, 255)
    p_b.alignment = PP_ALIGN.CENTER

    # 3. Left Column: Key System Capabilities Card
    cap_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.5), Inches(6.0), Inches(2.7))
    cap_card.fill.solid()
    cap_card.fill.fore_color.rgb = CARD_BG
    cap_card.line.color.rgb = RGBColor(51, 65, 85)
    
    tf_cap = cap_card.text_frame
    tf_cap.word_wrap = True
    p_cap_head = tf_cap.paragraphs[0]
    p_cap_head.text = "KEY SYSTEM CAPABILITIES"
    p_cap_head.font.name = 'Arial'
    p_cap_head.font.bold = True
    p_cap_head.font.size = Pt(13)
    p_cap_head.font.color.rgb = ACCENT_CYAN

    capabilities = [
        ("Intrusion Detection", "Real-time detection of human & vehicle perimeter breaches under zero light."),
        ("Thermal Integration", "Multi-spectral thermal camera feeds for all-weather 24/7 target identification."),
        ("Low Latency Alerts", "Sub-second threat response triggers directly to Command & Control posts."),
        ("Edge AI Processing", "On-device inference minimizing bandwidth usage across remote border zones.")
    ]

    for title, desc in capabilities:
        p_t = tf_cap.add_paragraph()
        p_t.text = f"• {title}: {desc}"
        p_t.font.name = 'Arial'
        p_t.font.size = Pt(10.5)
        p_t.font.color.rgb = TEXT_WHITE

    # 4. Left Column: Operational Advantages Card
    adv_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(4.4), Inches(6.0), Inches(2.5))
    adv_card.fill.solid()
    adv_card.fill.fore_color.rgb = CARD_BG
    adv_card.line.color.rgb = RGBColor(51, 65, 85)
    
    tf_adv = adv_card.text_frame
    tf_adv.word_wrap = True
    p_adv_head = tf_adv.paragraphs[0]
    p_adv_head.text = "OPERATIONAL ADVANTAGES"
    p_adv_head.font.name = 'Arial'
    p_adv_head.font.bold = True
    p_adv_head.font.size = Pt(13)
    p_adv_head.font.color.rgb = ACCENT_GREEN

    advantages = [
        ("Zero Fatigue Gap", "Eliminates the 95% loss in operator vigilance typically seen after 20 minutes of manual monitoring."),
        ("False Alarm Reduction", "AI filtering isolates real threats from foliage, wildlife, and ambient weather noise."),
        ("Scalable Infrastructure", "Seamlessly retrofits existing legacy CCTV and long-range thermal sensors.")
    ]

    for title, desc in advantages:
        p_a = tf_adv.add_paragraph()
        p_a.text = f"✔ {title}: {desc}"
        p_a.font.name = 'Arial'
        p_a.font.size = Pt(10.5)
        p_a.font.color.rgb = TEXT_WHITE

    # 5. Right Column: Performance Transformation Card
    perf_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.5), Inches(5.8), Inches(3.2))
    perf_card.fill.solid()
    perf_card.fill.fore_color.rgb = CARD_BG
    perf_card.line.color.rgb = RGBColor(51, 65, 85)

    tf_perf = perf_card.text_frame
    tf_perf.word_wrap = True
    p_perf_head = tf_perf.paragraphs[0]
    p_perf_head.text = "PERFORMANCE TRANSFORMATION (BEFORE vs AFTER)"
    p_perf_head.font.name = 'Arial'
    p_perf_head.font.bold = True
    p_perf_head.font.size = Pt(12)
    p_perf_head.font.color.rgb = ACCENT_CYAN

    metrics_list = [
        ("Threat Detection Accuracy", "35%  ➔  98%"),
        ("24/7 Active Monitoring", "15%  ➔  100%"),
        ("Response Efficiency", "30%  ➔  95%")
    ]

    for label, change in metrics_list:
        p_m = tf_perf.add_paragraph()
        p_m.text = f"\n• {label}: {change}"
        p_m.font.name = 'Arial'
        p_m.font.size = Pt(12)
        p_m.font.bold = True
        p_m.font.color.rgb = TEXT_WHITE

    # 6. Right Column Bottom: Key Statistics Boxes
    stats_data = [
        ("< 1 sec", "Alert Latency"),
        ("99.2%", "Precision Rate"),
        ("24/7/365", "Continuous Coverage")
    ]

    for i, (val, lbl) in enumerate(stats_data):
        box_left = Inches(6.9 + i * 1.98)
        s_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, box_left, Inches(4.9), Inches(1.84), Inches(2.0))
        s_box.fill.solid()
        s_box.fill.fore_color.rgb = INNER_CARD_BG
        s_box.line.color.rgb = ACCENT_CYAN

        tf_s = s_box.text_frame
        tf_s.word_wrap = True
        
        p_val = tf_s.paragraphs[0]
        p_val.text = val
        p_val.font.name = 'Arial'
        p_val.font.size = Pt(18)
        p_val.font.bold = True
        p_val.font.color.rgb = ACCENT_CYAN
        p_val.alignment = PP_ALIGN.CENTER

        p_lbl = tf_s.add_paragraph()
        p_lbl.text = lbl
        p_lbl.font.name = 'Arial'
        p_lbl.font.size = Pt(10)
        p_lbl.font.color.rgb = TEXT_MUTED
        p_lbl.alignment = PP_ALIGN.CENTER

    # Save output file
    output_filename = "IBVAP_Presentation.pptx"
    prs.save(output_filename)
    print(f"Presentation successfully generated: {output_filename}")

if __name__ == "__main__":
    create_presentation()